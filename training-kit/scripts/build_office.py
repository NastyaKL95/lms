#!/usr/bin/env python3
"""Собирает презентации PowerPoint и методические материалы Word."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn as ppt_qn
from pptx.util import Emu, Inches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[1]
PPTX_DIR = ROOT / "pptx"
DOCX_DIR = ROOT / "docx"

FONT = "Calibri"
COURSE = "Основы проектирования цифровых образовательных ресурсов"

NAVY = PptRGB(0x0F, 0x17, 0x2A)
BLUE = PptRGB(0x25, 0x63, 0xEB)
TEAL = PptRGB(0x0F, 0x76, 0x6E)
AMBER = PptRGB(0xB4, 0x53, 0x09)
BG = PptRGB(0xF8, 0xFA, 0xFC)
WHITE = PptRGB(0xFF, 0xFF, 0xFF)
INK = PptRGB(0x0F, 0x17, 0x2A)
MUTED = PptRGB(0x47, 0x55, 0x69)
LINE = PptRGB(0xE2, 0xE8, 0xF0)
SOFT = PptRGB(0xEF, 0xF6, 0xFF)
SOFT_TEAL = PptRGB(0xF0, 0xFD, 0xFA)
SOFT_AMBER = PptRGB(0xFF, 0xFB, 0xEB)
SLATE = PptRGB(0xCB, 0xD5, 0xE1)

W = 13.333
H = 7.5


def rgb(value: PptRGB) -> RGBColor:
    return RGBColor(value[0], value[1], value[2])


def _set_typeface(run, name: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        node = r_pr.find(ppt_qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, ppt_qn(tag))
        node.set("typeface", name)


def write_tf(tf, paragraphs, *, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor="t", spacing=1.05, margin=(0.1, 0.1, 0.06, 0.06)) -> None:
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(margin[0])
    tf.margin_right = Inches(margin[1])
    tf.margin_top = Inches(margin[2])
    tf.margin_bottom = Inches(margin[3])
    tf._txBody.bodyPr.set("anchor", anchor)
    for index, paragraph in enumerate(paragraphs):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = PptPt(0)
        p.space_after = PptPt(0)
        p.line_spacing = spacing
        runs = paragraph if isinstance(paragraph, list) else [(paragraph, bold, color, size)]
        for text, run_bold, run_color, run_size in runs:
            run = p.add_run()
            run.text = text
            run.font.size = PptPt(run_size)
            run.font.bold = run_bold
            run.font.color.rgb = run_color
            run.font.name = FONT
            _set_typeface(run, FONT)


def add_text(slide, x, y, w, h, paragraphs, **kwargs):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write_tf(shape.text_frame, paragraphs, **kwargs)
    return shape


def _no_line(shape) -> None:
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill, line=None, line_w=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        _no_line(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = PptPt(line_w)
    return shape


def rrect(slide, x, y, w, h, fill, line=None, line_w=1.0, radius=0.12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        _no_line(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = PptPt(line_w)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_chrome(slide, eyebrow, title, page, total):
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0, 0, W, 0.08, BLUE)
    add_text(slide, 0.48, 0.2, 10.5, 0.28, [eyebrow], size=12, bold=True, color=TEAL, margin=(0, 0, 0, 0))
    add_text(slide, 0.46, 0.46, 12.4, 0.62, [title], size=28, bold=True, color=NAVY, margin=(0, 0, 0, 0))
    rect(slide, 0, 7.22, W, 0.28, WHITE)
    rect(slide, 0, 7.22, W, 0.012, LINE)
    add_text(
        slide, 0.48, 7.24, 8.2, 0.24, [COURSE],
        size=11, color=MUTED, anchor="ctr", margin=(0, 0, 0, 0),
    )
    add_text(
        slide, 10.3, 7.24, 2.55, 0.24, [f"{page:02d}  /  {total:02d}"],
        size=11, bold=True, color=MUTED, align=PP_ALIGN.RIGHT, anchor="ctr", margin=(0, 0, 0, 0),
    )


def render_title(slide, spec):
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, 0, 0.16, H, TEAL)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.4), Inches(-1.4), Inches(4.6), Inches(4.6))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    _no_line(circle)
    srgb = circle.fill._xPr.find(ppt_qn("a:solidFill")).find(ppt_qn("a:srgbClr"))
    alpha = etree.SubElement(srgb, ppt_qn("a:alpha"))
    alpha.set("val", "18000")

    add_text(slide, 0.7, 0.48, 10, 0.3, [spec["kicker"]], size=14, bold=True, color=PptRGB(0x5E, 0xEA, 0xD4), margin=(0, 0, 0, 0))
    rect(slide, 0.72, 0.92, 1.5, 0.035, TEAL)
    add_text(
        slide, 0.68, 1.2, 11.6, 1.85, spec["title"].split("\n"),
        size=36, bold=True, color=WHITE, margin=(0, 0, 0, 0), spacing=0.95,
    )
    add_text(slide, 0.7, 3.15, 10.8, 1.05, [spec["subtitle"]], size=20, color=SLATE, margin=(0, 0, 0, 0))

    x = 0.72
    for label in spec["chips"]:
        width = max(1.7, min(4.4, len(label) * 0.125 + 0.5))
        chip = rrect(slide, x, 4.55, width, 0.42, PptRGB(0x1E, 0x29, 0x3B), radius=0.5)
        write_tf(chip.text_frame, [label], size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor="ctr", margin=(0.08, 0.08, 0, 0))
        x += width + 0.14

    add_text(slide, 0.7, 6.55, 11.5, 0.4, [spec["footer"]], size=14, color=SLATE, margin=(0, 0, 0, 0))


def render_cards(slide, spec):
    items = spec["items"]
    n = len(items)
    cols = 1 if n == 1 else 2 if n <= 4 else 3
    rows = (n + cols - 1) // cols
    left, top = 0.45, 1.32
    gap_x, gap_y = 0.18, 0.16
    grid_w, grid_h = 12.42, 5.7
    card_w = (grid_w - gap_x * (cols - 1)) / cols
    card_h = (grid_h - gap_y * (rows - 1)) / rows
    for index, text in enumerate(items):
        row, col = divmod(index, cols)
        items_in_row = min(cols, n - row * cols)
        row_width = items_in_row * card_w + (items_in_row - 1) * gap_x
        row_left = left + (grid_w - row_width) / 2
        x = row_left + col * (card_w + gap_x)
        y = top + row * (card_h + gap_y)
        card = rrect(slide, x, y, card_w, card_h, WHITE, LINE, 1.0, 0.08)
        rect(slide, x, y, card_w, 0.08, BLUE if index % 2 == 0 else TEAL)
        write_tf(
            card.text_frame,
            [
                [(f"{index + 1:02d}", True, TEAL, 18)],
                [(text, False, INK, 18 if card_w > 5 else 16)],
            ],
            anchor="ctr",
            margin=(0.28, 0.24, 0.16, 0.16),
            spacing=1.05,
        )


def _item_paragraph(item, size):
    if isinstance(item, tuple):
        label, text = item
        return [(label, True, BLUE, size), (text, False, INK, size)]
    return [(item, False, INK, size)]


def render_rows(slide, spec):
    items = spec["items"]
    lead = spec.get("lead")
    note = spec.get("note")
    top = 1.28
    bottom = 6.28 if note else 7.02
    if lead:
        lead_h = 0.78 if len(lead) > 110 else 0.58
        add_text(slide, 0.48, top, 12.3, lead_h, [lead], size=16, color=MUTED, margin=(0, 0, 0, 0))
        top += lead_h + 0.08
    n = len(items)
    gap = 0.1
    avail = bottom - top
    height = min(0.96, (avail - gap * (n - 1)) / n)
    size = 16 if height >= 0.72 else 14
    for index, item in enumerate(items):
        y = top + index * (height + gap)
        rrect(slide, 0.45, y, 12.42, height, WHITE, LINE, 1.0, 0.15)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.64), Inches(y + height / 2 - 0.21), Inches(0.42), Inches(0.42)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLUE
        _no_line(badge)
        write_tf(
            badge.text_frame, [str(index + 1)],
            size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor="ctr", margin=(0, 0, 0, 0),
        )
        box = add_text(
            slide, 1.2, y, 11.4, height, [_item_paragraph(item, size)],
            anchor="ctr", margin=(0.08, 0.16, 0.04, 0.04),
        )
        box  # text box sits above the card
    if note:
        bar = rrect(slide, 0.45, 6.42, 12.42, 0.66, SOFT, radius=0.12)
        write_tf(bar.text_frame, [note], size=15, bold=True, color=NAVY, anchor="ctr", margin=(0.22, 0.18, 0.04, 0.04))


def render_steps(slide, spec):
    items = spec["items"]
    n = len(items)
    top, bottom = 1.35, 7.02
    gap = 0.12
    height = min(0.92, (bottom - top - gap * (n - 1)) / n)
    for index, text in enumerate(items):
        y = top + index * (height + gap)
        rrect(slide, 0.45, y, 12.42, height, WHITE, LINE, 1.0, 0.12)
        rect(slide, 0.45, y, 0.1, height, TEAL)
        tag = rrect(slide, 0.74, y + height / 2 - 0.2, 1.15, 0.4, SOFT_TEAL, radius=0.4)
        write_tf(tag.text_frame, [f"{index + 1:02d}"], size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER, anchor="ctr", margin=(0, 0, 0, 0))
        add_text(slide, 2.05, y, 10.5, height, [text], size=16, anchor="ctr", margin=(0.08, 0.16, 0, 0))


def set_cell_border(cell, color="E2E8F0", width=6350):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        tag = ppt_qn(f"a:{edge}")
        for child in list(tc_pr.findall(tag)):
            tc_pr.remove(child)
        ln = etree.SubElement(tc_pr, tag)
        ln.set("w", str(width))
        solid = etree.SubElement(ln, ppt_qn("a:solidFill"))
        srgb = etree.SubElement(solid, ppt_qn("a:srgbClr"))
        srgb.set("val", color)


def style_ppt_cell(cell, text, *, fill, color, size, bold=False, align=PP_ALIGN.LEFT):
    cell.text = ""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = 3  # middle
    try:
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
    except Exception:
        pass
    write_tf(cell.text_frame, [text], size=size, bold=bold, color=color, align=align, anchor="ctr", margin=(0.08, 0.08, 0.02, 0.02))
    set_cell_border(cell)


def render_table(slide, spec):
    headers = spec["headers"]
    rows = spec["rows"]
    note = spec.get("note")
    widths = spec["widths"]
    top = 1.38
    table_h = 5.15 if note else 5.55
    row_h = table_h / (len(rows) + 1)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.48), Inches(top), Inches(sum(widths)), Inches(table_h))
    table = shape.table
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    for idx, header in enumerate(headers):
        style_ppt_cell(table.cell(0, idx), header, fill=NAVY, color=WHITE, size=14, bold=True)
    for r, row in enumerate(rows, start=1):
        bg = WHITE if r % 2 else PptRGB(0xF1, 0xF5, 0xF9)
        for c, value in enumerate(row):
            weight = c == 0
            style_ppt_cell(
                table.cell(r, c), value,
                fill=bg, color=NAVY if weight else INK, size=14, bold=weight,
            )
    for row in table.rows:
        row.height = Inches(row_h)
    if note:
        add_text(slide, 0.5, 6.62, 12.3, 0.48, [note], size=14, color=MUTED, margin=(0, 0, 0, 0))


def render_split(slide, spec):
    left = rrect(slide, 0.42, 1.32, 7.55, 5.7, WHITE, LINE, 1.0, 0.06)
    rect(slide, 0.42, 1.32, 7.55, 0.08, BLUE)
    write_tf(
        left.text_frame,
        [[(spec["left_title"], True, NAVY, 18)]] + [[(f"•  {item}", False, INK, 16)] for item in spec["left_items"]],
        anchor="t",
        margin=(0.28, 0.24, 0.28, 0.2),
        spacing=1.15,
    )
    right = rrect(slide, 8.15, 1.32, 4.75, 5.7, NAVY, radius=0.06)
    paragraphs = [[(spec["right_title"], True, PptRGB(0x5E, 0xEA, 0xD4), 16)]]
    paragraphs.append([(spec["right_text"], False, WHITE, 18)])
    write_tf(right.text_frame, paragraphs, anchor="ctr", margin=(0.32, 0.28, 0.2, 0.2), spacing=1.12)


def render_case(slide, spec):
    left = rrect(slide, 0.42, 1.32, 7.7, 5.7, WHITE, LINE, 1.0, 0.06)
    write_tf(
        left.text_frame,
        [
            [("СИТУАЦИЯ", True, AMBER, 13)],
            [(spec["situation"], False, NAVY, 22)],
        ],
        anchor="ctr",
        margin=(0.38, 0.36, 0.3, 0.3),
        spacing=1.12,
    )
    right = rrect(slide, 8.3, 1.32, 4.6, 5.7, NAVY, radius=0.06)
    paragraphs = [[("ЗАДАНИЕ", True, PptRGB(0x5E, 0xEA, 0xD4), 13)]]
    for index, task in enumerate(spec["tasks"], start=1):
        paragraphs.append([(f"{index}.  {task}", False, WHITE, 16)])
    write_tf(right.text_frame, paragraphs, anchor="ctr", margin=(0.32, 0.28, 0.24, 0.24), spacing=1.15)


def render_takeaways(slide, spec):
    items = spec["items"]
    gap = 0.18
    width = (12.42 - gap * (len(items) - 1)) / len(items)
    for index, text in enumerate(items):
        x = 0.45 + index * (width + gap)
        card = rrect(slide, x, 1.7, width, 4.7, WHITE, LINE, 1.0, 0.08)
        rect(slide, x, 1.7, width, 0.1, TEAL if index != 1 else BLUE)
        write_tf(
            card.text_frame,
            [[(f"{index + 1:02d}", True, TEAL, 22)], [(text, False, NAVY, 18)]],
            anchor="ctr",
            margin=(0.28, 0.24, 0.2, 0.2),
        )


def render_deck(path: Path, section: str, slides: list[dict], title: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = title
    prs.core_properties.subject = COURSE
    prs.core_properties.author = "Учебный комплект"
    prs.core_properties.category = "Презентация к занятию"
    try:
        prs.core_properties.language = "ru-RU"
    except Exception:
        pass
    total = len(slides)
    for index, spec in enumerate(slides, start=1):
        slide = blank(prs)
        layout = spec["layout"]
        if layout == "title":
            render_title(slide, spec)
            continue
        content_chrome(slide, section, spec["title"], index, total)
        {
            "cards": render_cards,
            "rows": render_rows,
            "steps": render_steps,
            "table": render_table,
            "split": render_split,
            "case": render_case,
            "takeaways": render_takeaways,
        }[layout](slide, spec)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def deck_one():
    section = "ТЕМА 1   ·   ЦОР И ПРАВОВЫЕ АСПЕКТЫ"
    slides = [
        {
            "layout": "title",
            "kicker": "МОДУЛЬ 1",
            "title": "Основные понятия\nи классификация ЦОР",
            "subtitle": "Нормативно-правовые аспекты разработки образовательного контента",
            "chips": ["Лекция", "Мини-кейсы", "Законодательство РФ"],
            "footer": "Педагогика   ·   технология   ·   право",
        },
        {
            "layout": "cards",
            "title": "Цели занятия",
            "items": [
                "Понять, что такое ЦОР и какую задачу они решают в обучении",
                "Выбрать формат ресурса под учебную цель, а не «по привычке»",
                "Применять правила авторского права, лицензий и защиты ПДн",
                "Проектировать ЦОР с учетом ФГОС, СанПиН и доступности",
            ],
        },
        {
            "layout": "rows",
            "title": "Что такое ЦОР",
            "lead": "Цифровой образовательный ресурс — это контент или инструмент, который помогает объяснять, тренировать, проверять и обсуждать учебный материал.",
            "items": [
                "Интерактивные уроки и микрокурсы",
                "Тесты и тренажеры",
                "Виртуальные лаборатории",
                "Видео- и аудиолекции",
                "Симуляторы профессиональных ситуаций",
            ],
        },
        {
            "layout": "cards",
            "title": "Зачем ЦОР в современном обучении",
            "items": [
                "Студент движется в своем темпе и по своей траектории",
                "Интерактив удерживает внимание лучше длинного монолога",
                "Преподаватель видит прогресс по данным, а не «на глаз»",
                "Один качественный ресурс можно использовать повторно",
                "Курс работает очно, дистанционно и в смешанном формате",
                "Результат измеряется заданием, а не только присутствием",
            ],
        },
        {
            "layout": "rows",
            "title": "Как классифицировать ЦОР",
            "items": [
                ("По назначению.  ", "объяснение, тренировка, контроль, рефлексия"),
                ("По интерактивности.  ", "пассивные, активные, адаптивные"),
                ("По формату.  ", "текст, видео, симуляция, игровая механика"),
                ("По уровню.  ", "базовый, продвинутый, экспертный"),
                ("По доступу.  ", "закрытые материалы и открытые ресурсы (OER)"),
            ],
        },
        {
            "layout": "rows",
            "title": "Форматы и стандарты",
            "items": [
                ("SCORM.  ", "курс открывается в LMS и передает результат прохождения"),
                ("xAPI.  ", "фиксирует действия обучающегося, а не только «просмотрено»"),
                ("LTI.  ", "подключает внешний инструмент внутрь LMS"),
                ("HTML5.  ", "работает на компьютере и телефоне без отдельного плеера"),
            ],
            "note": "Стандарт выбирают от учебной цели. Модный формат без методики не делает курс лучше.",
        },
        {
            "layout": "table",
            "title": "Правовой контур: что проверять всегда",
            "headers": ["Норма", "О чем помнить при разработке"],
            "widths": [4.7, 7.7],
            "rows": [
                ["ГК РФ, часть IV", "Авторские права и условия использования произведений"],
                ["273-ФЗ «Об образовании»", "Требования к организации образовательного процесса"],
                ["149-ФЗ «Об информации»", "Общие правила работы с информацией и ее защиты"],
                ["152-ФЗ «О персональных данных»", "Сбор, хранение и доступ к данным обучающихся"],
                ["181-ФЗ и доступность", "Инклюзия и доступная образовательная среда"],
                ["СанПиН", "Гигиена цифровой нагрузки и режим занятий"],
            ],
            "note": "Перед публикацией сверяйте актуальные редакции и локальные акты своей организации.",
        },
        {
            "layout": "rows",
            "title": "Авторское право: типичные ошибки",
            "items": [
                "Картинка из поиска без лицензии и указания автора",
                "Фрагмент учебника без правового основания",
                "Курс отдан подрядчику без договора о правах",
                "У собственного материала не указаны условия использования",
            ],
            "note": "Последствия: претензия правообладателя, снятие материала, финансовые риски.",
        },
        {
            "layout": "steps",
            "title": "Как законно использовать чужой материал",
            "items": [
                "Найти источник и правообладателя",
                "Проверить лицензию, включая Creative Commons",
                "Убедиться, что можно перерабатывать и использовать в обучении",
                "Оформить атрибуцию: автор, источник, лицензия",
                "Записать основание в реестр материалов курса",
            ],
        },
        {
            "layout": "cards",
            "title": "Персональные данные обучающихся",
            "items": [
                "Сначала цель обработки, потом список полей",
                "Не собирать данные «на всякий случай»",
                "Иметь законное основание, включая согласие, где оно нужно",
                "Открывать данные только тем ролям, которым они нужны",
                "Защищать данные организационно и технически",
                "Закрепить порядок в локальных актах организации",
            ],
        },
        {
            "layout": "rows",
            "title": "Сверка с ФГОС и результатами обучения",
            "lead": "Каждый ресурс отвечает на четыре вопроса. Если ответа нет, ресурс еще не готов.",
            "items": [
                "Какой результат обучения он формирует?",
                "Какая компетенция или индикатор проверяется?",
                "Понятно ли студенту, что считается успехом?",
                "Где и как фиксируется результат в оценивании?",
            ],
        },
        {
            "layout": "rows",
            "title": "СанПиН и удобство для восприятия",
            "items": [
                "Экранная работа чередуется с обсуждением и практикой",
                "Длительность непрерывного просмотра разумная",
                "Текст крупный, контрастный и без визуального шума",
                "На слайде одна мысль, а не конспект параграфа",
                "В сценарии есть паузы и смена вида деятельности",
            ],
        },
        {
            "layout": "cards",
            "title": "Доступность для обучающихся с ОВЗ",
            "items": [
                "Субтитры и текстовая версия аудио и видео",
                "Курс можно пройти с клавиатуры, без точного клика",
                "Контраст текста и фона достаточный для чтения",
                "Структура простая: заголовки, короткие абзацы, понятные кнопки",
                "Материал совместим со средствами чтения с экрана",
            ],
        },
        {
            "layout": "case",
            "title": "Мини-кейс на 10 минут",
            "situation": "Преподаватель собрал курс из роликов YouTube, картинок из поиска и формы, где студент указывает ФИО и телефон.",
            "tasks": [
                "Найдите пять правовых и методических рисков",
                "Предложите, как исправить каждый риск",
                "Соберите чек-лист «перед публикацией курса»",
            ],
        },
        {
            "layout": "takeaways",
            "title": "Что уносим с занятия",
            "items": [
                "ЦОР работает, когда связаны цель, активность и проверка",
                "Правовая проверка — часть разработки, а не приложение в конце",
                "Красивый экран не заменяет результат обучения",
            ],
        },
        {
            "layout": "steps",
            "title": "Самостоятельная работа",
            "items": [
                "Сравните модели педагогического дизайна ADDIE, ASSURE и SAM",
                "Выберите одну учебную тему и набросайте сценарий ЦОР",
                "Составьте правовой чек-лист именно для этого ресурса",
                "Опишите, как будут обеспечены доступность и связь с ФГОС",
            ],
        },
    ]
    render_deck(PPTX_DIR / "Тема 1. Основные понятия и классификация ЦОР.pptx", section, slides, "Тема 1. ЦОР и правовые аспекты")


def deck_two():
    section = "ТЕМА 2   ·   EDTECH И ПЛАТФОРМЫ"
    slides = [
        {
            "layout": "title",
            "kicker": "МОДУЛЬ 1",
            "title": "Современные тенденции\nв EdTech",
            "subtitle": "Как сравнивать платформы и конструкторы курсов, а не выбирать «самую известную»",
            "chips": ["Лекция", "Сравнительный анализ", "Подготовка к практике"],
            "footer": "Цель курса   ·   критерии   ·   пилот",
        },
        {
            "layout": "cards",
            "title": "Цели занятия",
            "items": [
                "Назвать рабочие тренды EdTech и отделить их от лозунгов",
                "Сопоставить российские и зарубежные платформы",
                "Собрать критерии выбора LMS и конструктора",
                "Подготовить сравнение для практического занятия",
            ],
        },
        {
            "layout": "rows",
            "title": "Что сейчас меняет цифровое обучение",
            "items": [
                "Адаптивные траектории вместо одного маршрута для всех",
                "ИИ-помощники для черновиков, заданий и обратной связи",
                "Короткие модули вместо длинных неделимых курсов",
                "Аналитика обучения: где группа реально теряет тему",
                "Совместная работа и обсуждение, а не только просмотр",
            ],
        },
        {
            "layout": "split",
            "title": "Искусственный интеллект в курсе",
            "left_title": "Где ИИ полезен",
            "left_items": [
                "Черновик объяснения, который преподаватель правит",
                "Несколько вариантов одного задания",
                "Проверка типовых ответов с понятным ключом",
                "Подсказка, какой теме стоит вернуться",
                "Рекомендация следующего шага студенту",
            ],
            "right_title": "ГРАНИЦА",
            "right_text": "ИИ не заменяет педагогический дизайн и решение преподавателя. Оценку, которая влияет на аттестацию, подтверждает человек.",
        },
        {
            "layout": "steps",
            "title": "Из чего состоит адаптивное обучение",
            "items": [
                "Входная диагностика: что студент уже умеет",
                "Ветвление: разный следующий шаг для разного результата",
                "Сложность меняется по ходу, а не только между курсами",
                "Короткая обратная связь сразу после ошибки",
                "Преподаватель видит, почему система предложила этот путь",
            ],
        },
        {
            "layout": "rows",
            "title": "Геймификация, которая учит",
            "items": [
                "Прогресс показывает движение по навыку, а не только баллы",
                "Знак отличия выдается за конкретное умение",
                "Командное задание требует вклада каждого",
                "Короткий кейс со сроком имитирует рабочее ограничение",
                "Миссия связана с реальной профессиональной задачей",
            ],
            "note": "Игра усиливает учебную цель. Если цель исчезла, осталась только оболочка.",
        },
        {
            "layout": "table",
            "title": "Платформы: сильные стороны и цена выбора",
            "headers": ["Платформа", "Сильные стороны", "На что смотреть внимательно"],
            "widths": [2.5, 5.0, 4.9],
            "rows": [
                ["Moodle", "Гибкость, открытый код, большое сообщество", "Нужны настройка и администратор"],
                ["Canvas", "Современный интерфейс и интеграции", "Стоимость владения часто выше"],
                ["Blackboard", "Зрелые корпоративные функции", "Внедрение может быть сложным"],
                ["Stepik", "Удобно собирать онлайн-курс", "Меньше свободы классической LMS"],
                ["ЯКласс", "Готовые материалы, школьный фокус", "Логика платформы задает рамку"],
            ],
            "note": "Платформа хороша в конкретном курсе и команде. Универсального победителя нет.",
        },
        {
            "layout": "table",
            "title": "Конструкторы курсов",
            "headers": ["Инструмент", "Когда он уместен"],
            "widths": [3.6, 8.8],
            "rows": [
                ["iSpring", "Нужно быстро собрать курс из уже готовой презентации"],
                ["Articulate 360", "Нужны ветвления, диалоговые тренажеры и сценарии"],
                ["H5P", "Нужны легкие интерактивы прямо внутри LMS"],
                ["Canva и аналоги", "Нужны наглядные схемы, карточки и визуальные конспекты"],
            ],
            "note": "Критерий выбора: как быстро команда получает качество, понятное именно вашей аудитории.",
        },
        {
            "layout": "steps",
            "title": "Семь критериев сравнения",
            "items": [
                "Соответствие целям обучения, а не списку функций",
                "Удобство и для преподавателя, и для студента",
                "Интеграции: LMS, вход в систему, видео, аналитика",
                "Отчеты, по которым видно прогресс и трудные места",
                "Стоимость запуска и поддержки, а не только лицензии",
                "Надежность при росте числа курсов и пользователей",
                "Безопасность и соответствие требованиям РФ",
            ],
        },
        {
            "layout": "rows",
            "title": "Право и этика при выборе сервиса",
            "items": [
                "Понятно, где хранятся персональные данные и кто к ним допущен",
                "Студенту объясняют, где в курсе используется ИИ",
                "Автоматическая оценка не становится единственным решением",
                "Права на контент, созданный в платформе, зафиксированы",
                "Сервис не противоречит локальным актам организации",
            ],
        },
        {
            "layout": "steps",
            "title": "Как выбрать платформу без спешки",
            "items": [
                "Записать цели курса и ограничения: люди, сроки, бюджет, нормы",
                "Назначить вес каждому критерию до просмотра рекламы",
                "Провести пилот на одном-двух модулях, а не на всем курсе",
                "Собрать отзыв студентов и преподавателей по одним вопросам",
                "Принять решение и составить короткий план внедрения",
            ],
        },
        {
            "layout": "case",
            "title": "Задача для группы",
            "situation": "Нужно запустить курс для вашей аудитории. Времени на «идеальную платформу» нет, но ошибка выбора будет дорогой.",
            "tasks": [
                "Возьмите 2 LMS и 1 конструктор",
                "Сравните их по семи критериям",
                "Обоснуйте выбор под конкретный курс",
                "Назовите риски, которые надо закрыть до старта",
            ],
        },
        {
            "layout": "takeaways",
            "title": "Выводы",
            "items": [
                "Лучшая платформа — та, что тянет вашу методику",
                "Решают методика, готовность команды и поддержка",
                "Пилот полезнее самой убедительной презентации вендора",
            ],
        },
        {
            "layout": "steps",
            "title": "Самостоятельная работа",
            "items": [
                "Сравните 3 платформы и 2 конструктора",
                "Оцените их минимум по семи критериям",
                "Сформулируйте итоговый выбор и три аргумента",
                "Добавьте риски и меры, которыми эти риски снижаются",
            ],
        },
    ]
    render_deck(PPTX_DIR / "Тема 2. Современные тенденции в EdTech.pptx", section, slides, "Тема 2. EdTech и платформы")


def deck_three():
    section = "ТЕМА 3   ·   LMS MOODLE"
    slides = [
        {
            "layout": "title",
            "kicker": "МОДУЛЬ 1",
            "title": "Системы управления\nобучением",
            "subtitle": "Собираем рабочий курс в Moodle: структура, роли, задания и проверка",
            "chips": ["Лекция", "Практика 1–3", "Мини-проект"],
            "footer": "Структура   ·   роли   ·   оценивание",
        },
        {
            "layout": "cards",
            "title": "Цели занятия",
            "items": [
                "Понять, за что отвечает LMS, а за что — преподаватель",
                "Собрать понятную структуру курса",
                "Настроить роли, доступ и правила оценивания",
                "Довести тестовый курс до пилотного запуска",
            ],
        },
        {
            "layout": "split",
            "title": "Что дает LMS",
            "left_title": "Среда, в которой живет курс",
            "left_items": [
                "Материалы лежат в одном месте и в нужном порядке",
                "Есть сроки, задания и единые правила сдачи",
                "Обсуждение не теряется в личных переписках",
                "Результаты видны по человеку и по группе",
            ],
            "right_title": "ГЛАВНАЯ ЦЕННОСТЬ",
            "right_text": "Курс можно повторить с новой группой без сборки «с нуля» и без потери правил оценивания.",
        },
        {
            "layout": "steps",
            "title": "Архитектура курса в Moodle",
            "items": [
                "Категория курсов — полка, на которой лежит программа",
                "Курс — цельная учебная единица со своими целями",
                "Разделы — темы или недели, по которым идет студент",
                "Элементы — лекция, задание, тест, форум, файл",
                "Оценивание — журнал, критерии и условия завершения",
            ],
        },
        {
            "layout": "rows",
            "title": "Роли: кому что можно",
            "items": [
                ("Администратор.  ", "системные настройки площадки"),
                ("Менеджер курса.  ", "структура курса и состав участников"),
                ("Преподаватель.  ", "материалы, проверка и обратная связь"),
                ("Студент.  ", "изучение, сдача работ и обсуждение"),
                ("Гость.  ", "только просмотр, и только если это включено"),
            ],
            "note": "Правило безопасности: выдавать минимальные права, которых хватает для работы.",
        },
        {
            "layout": "steps",
            "title": "Пять шагов до запуска",
            "items": [
                "Создать курс и прямо на старте написать цели и результаты",
                "Разбить содержание на короткие тематические модули",
                "В каждый модуль добавить материал и действие студента",
                "Поставить сроки и условия, при которых тема считается пройденной",
                "Пройти курс под учетной записью студента до публикации",
            ],
        },
        {
            "layout": "rows",
            "title": "Какой элемент для какой задачи",
            "items": [
                ("Страница или лекция.  ", "объяснить теорию короткими шагами"),
                ("Задание.  ", "принять работу и дать развернутый комментарий"),
                ("Тест.  ", "проверить понимание и дать самопроверку"),
                ("Форум.  ", "организовать обсуждение или взаимную проверку"),
                ("Глоссарий или база.  ", "собирать общие примеры и термины"),
            ],
        },
        {
            "layout": "cards",
            "title": "Как настроить тест, чтобы он учил",
            "items": [
                "Вопросы разложены по темам, а не одной кучей",
                "Студент получает случайную выборку из банка",
                "Время ограничено и сказано заранее",
                "Можно несколько попыток, и после каждой есть подсказка",
                "Варианты ответов перемешиваются",
                "Сложный вопрос не решает судьбу всей темы",
            ],
        },
        {
            "layout": "rows",
            "title": "Что можно подключить к Moodle",
            "items": [
                "Видеосервисы для лекций и записей",
                "Вебинары для синхронных встреч",
                "Внешние интерактивы через LTI и H5P",
                "Проверку заимствований в письменных работах",
                "Корпоративный вход (SSO), чтобы не заводить лишние пароли",
            ],
        },
        {
            "layout": "cards",
            "title": "Какие данные смотреть каждую неделю",
            "items": [
                "Кто перестал заходить",
                "Где остановился прогресс по модулям",
                "Какие вопросы теста проваливает вся группа",
                "Есть ли живое обсуждение или только тишина",
                "Чем динамика одного студента отличается от группы",
            ],
        },
        {
            "layout": "rows",
            "title": "Ошибки, из-за которых курс «не взлетает»",
            "items": [
                "Первый модуль перегружен всеми файлами сразу",
                "Критерии оценки спрятаны или написаны слишком общо",
                "Много текста и почти нет действия студента",
                "Непонятно, куда нажимать и что делать дальше",
                "Работа сдана, а обратная связь приходит слишком поздно",
            ],
        },
        {
            "layout": "rows",
            "title": "Правовой минимум для курса в LMS",
            "items": [
                "Пользователь понимает, какие данные обрабатываются и зачем",
                "Доступ к персональным данным разграничен по ролям",
                "В курсе только те материалы, которые можно использовать",
                "Оценки хранятся по регламенту организации, а не «как получится»",
                "Сервер и резервные копии соответствуют требованиям ИБ",
            ],
        },
        {
            "layout": "case",
            "title": "Мини-проект: тестовый курс",
            "situation": "Соберите курс, который можно отдать пилотной группе уже после этого цикла практических занятий.",
            "tasks": [
                "Три тематических раздела с понятными целями",
                "Тест: банк не меньше 15 вопросов",
                "Одно задание с критериями оценки",
                "Форум и условия завершения разделов",
            ],
        },
        {
            "layout": "cards",
            "title": "По чему принимаем тестовый курс",
            "items": [
                "Структура читается без устных пояснений",
                "Цель, активность и оценка связаны между собой",
                "Журнал и критерии реально работают",
                "Инструкция студенту написана простым языком",
                "Контент опубликован правомерно",
            ],
        },
        {
            "layout": "steps",
            "title": "Самостоятельная доработка",
            "items": [
                "Добавьте развилку: разный шаг для разного результата",
                "Упростите путь студента по замечаниям с пилота",
                "Напишите короткий отчет: что изменили и зачем",
            ],
        },
    ]
    render_deck(PPTX_DIR / "Тема 3. Системы управления обучением LMS.pptx", section, slides, "Тема 3. LMS Moodle")


def deck_four():
    section = "ТЕМА 4   ·   ПРОЕКТИРОВАНИЕ ИС"
    slides = [
        {
            "layout": "title",
            "kicker": "МОДУЛЬ 1",
            "title": "Основы проектирования\nинформационных систем",
            "subtitle": "От обследования предметной области до диаграмм BPMN и UML",
            "chips": ["Лекции", "Обследование", "BPMN и UML"],
            "footer": "Процесс   ·   данные   ·   решение",
        },
        {
            "layout": "cards",
            "title": "Цели занятия",
            "items": [
                "Увидеть жизненный цикл системы целиком, а не только экран",
                "Различить функциональный и объектный подходы",
                "Выделить процессы, которые стоит автоматизировать",
                "Собрать базовые диаграммы BPMN и UML",
            ],
        },
        {
            "layout": "rows",
            "title": "Что значит спроектировать систему",
            "lead": "Проект — это согласованная модель будущей системы, по которой можно принимать решения до написания кода.",
            "items": [
                "Цели и границы: что входит в систему, а что нет",
                "Пользователи и роли",
                "Данные и бизнес-процессы",
                "Функции и требования к качеству",
                "Архитектурные решения и ограничения",
            ],
        },
        {
            "layout": "steps",
            "title": "Жизненный цикл информационной системы",
            "items": [
                "Инициация и обследование",
                "Сбор и анализ требований",
                "Проектирование",
                "Реализация",
                "Тестирование",
                "Внедрение",
                "Сопровождение и развитие",
            ],
        },
        {
            "layout": "cards",
            "title": "Какую модель цикла выбрать",
            "items": [
                "Каскад: этапы понятны, но поворот дорогой, если требования уже изменились",
                "Итерации: система улучшается циклами и каждый цикл дает рабочий кусок",
                "Спираль: в центре внимания риски и проверка самых опасных гипотез",
                "Гибкие подходы: короткие поставки и быстрая обратная связь",
            ],
        },
        {
            "layout": "split",
            "title": "Два подхода к проектированию",
            "left_title": "Функциональный",
            "left_items": [
                "Смотрим на функции и процессы",
                "Делим работу на входы и выходы",
                "Описываем регламент исполнения",
                "Ищем узкие места",
                "Инструменты: схемы, BPMN, матрица ролей",
            ],
            "right_title": "ОБЪЕКТНЫЙ",
            "right_text": "Смотрим на сущности предметной области: объекты, атрибуты, методы, состояния и связи. Инструмент — UML.",
        },
        {
            "layout": "rows",
            "title": "Системное проектирование связывает слои",
            "items": [
                "Бизнес-цели",
                "Процессы",
                "Данные",
                "Интерфейсы пользователей и внешних систем",
                "Технологические ограничения",
            ],
            "note": "Готовая модель непротиворечива: процесс, данные и экран рассказывают одну историю.",
        },
        {
            "layout": "cards",
            "title": "Что собрать на обследовании",
            "items": [
                "Короткое описание предметной области",
                "Участники и их роли",
                "Как процесс идет сейчас (As-Is)",
                "Где теряется время, возникают ошибки и дубли",
                "Ограничения: сроки, люди, нормы, действующие системы",
            ],
        },
        {
            "layout": "rows",
            "title": "Что стоит автоматизировать в первую очередь",
            "items": [
                "Операции, которые повторяются одинаково",
                "Шаги, где цена ошибки высокая",
                "Ручные согласования и перенос данных между людьми",
                "Отчеты, которые долго собирают вручную",
                "Решения, для которых уже есть понятные правила",
            ],
        },
        {
            "layout": "rows",
            "title": "BPMN: минимальный набор на диаграмме",
            "lead": "Диаграмма процесса отвечает на вопрос «кто что делает и в каком порядке».",
            "items": [
                "Участники: пулы и дорожки",
                "События старта и завершения",
                "Задачи и вложенные подпроцессы",
                "Шлюзы: где процесс расходится и сходится",
                "Потоки управления и сообщения между участниками",
            ],
        },
        {
            "layout": "split",
            "title": "UML: с чего начать",
            "left_title": "Диаграмма вариантов использования",
            "left_items": [
                "Кто пользователь системы",
                "Какую цель он закрывает",
                "Какие сценарии критичны",
                "Где проходит граница системы",
                "Старт: 5–7 главных сценариев, не 40",
            ],
            "right_title": "ДАЛЬШЕ",
            "right_text": "Sequence показывает порядок сообщений во времени. Activity показывает логику шагов и развилки. Детализируйте только то, что влияет на решение.",
        },
        {
            "layout": "rows",
            "title": "Где проекты обычно ломаются",
            "items": [
                "Схемы не показаны тем, кто реально выполняет процесс",
                "На старте нарисовано слишком много деталей",
                "Один термин означает разное у разных участников",
                "Забыли требования к скорости, надежности и безопасности",
                "Красивая схема не связана с тем, что будут программировать",
            ],
        },
        {
            "layout": "steps",
            "title": "Практический цикл",
            "items": [
                "Опишите выбранную предметную область и ее границы",
                "Составьте список задач автоматизации и расставьте приоритет",
                "Постройте BPMN текущего процесса",
                "Постройте UML: варианты использования и Sequence или Activity",
            ],
        },
        {
            "layout": "cards",
            "title": "Хорошая модель узнается так",
            "items": [
                "Ее понимает человек, который не рисовал диаграмму",
                "В ней есть главные сценарии, включая отказ и исключение",
                "В схемах нет взаимоисключающих шагов",
                "Каждая диаграмма отвечает на требование",
                "По модели уже можно ставить задачу на реализацию",
            ],
        },
        {
            "layout": "steps",
            "title": "Самостоятельный мини-проект",
            "items": [
                "Описание предметной области на одну-две страницы",
                "Приоритетный список задач автоматизации",
                "Комплект BPMN и UML по одному ключевому процессу",
                "Короткая защита: какое решение принято и почему",
            ],
        },
    ]
    render_deck(
        PPTX_DIR / "Тема 4. Основы проектирования информационных систем.pptx",
        section, slides, "Тема 4. Проектирование информационных систем",
    )


# --- Word ---

def set_run_font(run, name, size, bold=False, color=None):
    run.font.name = name
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, margin_dxa=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_dxa))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = int(sum(widths_cm) * 567)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    for width in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(width * 567)))
        grid.append(gc)
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_row_height(row, twips: int) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


def set_cell_borders(cell, color="D6DEE8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def cell_text(cell, text, *, size=11, bold=False, color=None, fill=None, align="left", center=False, border=True):
    if fill:
        shade_cell(cell, fill)
    if border:
        set_cell_borders(cell, "1E293B" if fill == "0F172A" else "D6DEE8")
    set_cell_margins(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, FONT, size, bold, color or RGBColor(0x0F, 0x17, 0x2A))
    if center:
        cell.vertical_alignment = 1


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, FONT, 9, color=RGBColor(0x47, 0x55, 0x69))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.15
    for style_name, size, color, before, after in (
        ("Heading 1", 18, RGBColor(0x0F, 0x17, 0x2A), 16, 8),
        ("Heading 2", 14, RGBColor(0x1D, 0x4E, 0xD8), 14, 6),
        ("Heading 3", 12, RGBColor(0x0F, 0x76, 0x6E), 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def new_doc(header_title: str) -> Document:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run(header_title)
    set_run_font(run, FONT, 9, color=RGBColor(0x47, 0x55, 0x69))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("Стр. ")
    set_run_font(label, FONT, 9, color=RGBColor(0x47, 0x55, 0x69))
    add_page_number(footer)
    doc.core_properties.author = "Учебный комплект"
    doc.core_properties.subject = COURSE
    doc.core_properties.language = "ru-RU"
    doc.core_properties.category = "Методические материалы"
    return doc


def add_cover(doc: Document, kicker: str, title: str, subtitle: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [17.4])
    cell = table.cell(0, 0)
    shade_cell(cell, "0F172A")
    set_cell_margins(cell, 160)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(kicker.upper())
    set_run_font(r1, FONT, 11, True, RGBColor(0x5E, 0xEA, 0xD4))
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(8)
    p2.paragraph_format.space_after = Pt(6)
    r2 = p2.add_run(title)
    set_run_font(r2, FONT, 22, True, RGBColor(0xFF, 0xFF, 0xFF))
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(2)
    r3 = p3.add_run(subtitle)
    set_run_font(r3, FONT, 12, False, RGBColor(0xCB, 0xD5, 0xE1))
    doc.add_paragraph()


def add_para(doc, text, *, size=12, bold=False, color=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, FONT, size, bold, color or RGBColor(0x0F, 0x17, 0x2A))
    return p


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.clear()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(0)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(lead)
            set_run_font(r1, FONT, 12, True, RGBColor(0x0F, 0x17, 0x2A))
            r2 = p.add_run(rest)
            set_run_font(r2, FONT, 12, False, RGBColor(0x0F, 0x17, 0x2A))
        else:
            run = p.add_run(item)
            set_run_font(run, FONT, 12, False, RGBColor(0x0F, 0x17, 0x2A))


def add_callout(doc, text, fill="EFF6FF"):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [17.4])
    cell = table.cell(0, 0)
    cell_text(cell, text, size=11, bold=False, fill=fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_data_table(doc, headers, rows, widths, header=True, row_twips=420):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for idx, title in enumerate(headers):
        cell_text(table.rows[0].cells[idx], title, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), fill="0F172A")
    for r, row in enumerate(rows):
        fill = "FFFFFF" if r % 2 == 0 else "F8FAFC"
        for c, value in enumerate(row):
            cell_text(table.rows[r + 1].cells[c], value, size=10, bold=(c == 0 and header), fill=fill)
        prevent_row_split(table.rows[r + 1])
        set_row_height(table.rows[r + 1], row_twips)
    prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_form_table(doc, rows_labels):
    table = doc.add_table(rows=len(rows_labels), cols=2)
    set_table_widths(table, [6.2, 11.2])
    for idx, label in enumerate(rows_labels):
        cell_text(table.rows[idx].cells[0], label, size=11, bold=True, fill="F8FAFC")
        cell_text(table.rows[idx].cells[1], " ", size=11, fill="FFFFFF")
        set_row_height(table.rows[idx], 900)
        prevent_row_split(table.rows[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def save_doc(doc: Document, name: str) -> None:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = name
    doc.save(DOCX_DIR / f"{name}.docx")


def build_overview():
    doc = new_doc(COURSE)
    add_cover(
        doc,
        "Комплект к модулю",
        "Материалы для проведения занятий",
        "Презентации PowerPoint и методические документы Word по четырем темам модуля.",
    )
    add_para(doc, "Комплект собран по темам рабочей программы «Основы проектирования цифровых образовательных ресурсов». В правовых блоках опора — законодательство РФ. Перед занятием проверяйте актуальные редакции норм и локальные акты организации.")
    doc.add_heading("Презентации", level=1)
    add_data_table(
        doc,
        ["Файл", "Когда открывать"],
        [
            ["Тема 1. Основные понятия и классификация ЦОР.pptx", "Лекция о типах ЦОР, стандартах и правовом контуре"],
            ["Тема 2. Современные тенденции в EdTech.pptx", "Лекция о трендах, LMS и конструкторах"],
            ["Тема 3. Системы управления обучением LMS.pptx", "Лекция и практика по Moodle"],
            ["Тема 4. Основы проектирования информационных систем.pptx", "Лекции и практика по обследованию, BPMN и UML"],
        ],
        [8.2, 9.2],
    )
    doc.add_heading("Документы для преподавателя и студентов", level=1)
    add_data_table(
        doc,
        ["Файл", "Назначение"],
        [
            ["Методический сценарий проведения занятий.docx", "Ход пары на 90 минут и подготовка"],
            ["Рабочая тетрадь.docx", "Бланки, которые студенты заполняют на практике"],
            ["Банк оценочных средств.docx", "Вопросы, ключи и рубрика"],
            ["Правовая памятка (РФ).docx", "Нормы и чек-лист перед публикацией"],
            ["Визуальный гайд презентаций.docx", "Как оформлять собственные слайды студентов"],
        ],
        [8.2, 9.2],
    )
    doc.add_heading("Как провести модуль", level=1)
    add_bullets(doc, [
        "На лекции показывайте презентацию темы и оставляйте время на кейс.",
        "На практике раздайте рабочую тетрадь: студенты сдают заполненные бланки.",
        "Текущий контроль берите из банка оценочных средств.",
        "В темах 1–3 держите под рукой правовую памятку.",
        "Зачет складывается из мини-проекта и короткой защиты.",
    ], numbered=True)
    save_doc(doc, "Содержание комплекта")


def build_instructor():
    doc = new_doc("Методический сценарий")
    add_cover(doc, "Для преподавателя", "Сценарий проведения занятий", "Модель пары на 90 минут и отдельный ход по каждой теме модуля.")
    doc.add_heading("Общая модель пары", level=1)
    add_data_table(
        doc,
        ["Этап", "Минуты", "Что происходит"],
        [
            ["Вход в тему", "10", "Вопрос к опыту группы и цели занятия"],
            ["Теория", "25", "Презентация: одна мысль на слайд, без чтения текста"],
            ["Кейс", "15", "Разбор ситуации с доски или из презентации"],
            ["Практика", "25", "Группы заполняют бланк рабочей тетради"],
            ["Рефлексия", "10", "Что получилось и какой риск остался"],
            ["Домашняя работа", "5", "Последний слайд: продукт и срок"],
        ],
        [4.2, 2.4, 10.8],
    )
    doc.add_heading("Тема 1. ЦОР и право", level=1)
    add_bullets(doc, [
        ("Старт. ", "Спросите, какими цифровыми ресурсами студенты пользуются каждую неделю."),
        ("Фокус. ", "Цепочка «учебная цель → формат ЦОР → правовое ограничение»."),
        ("Активность. ", "Кейс с чужим видео, картинками из поиска и формой с телефоном."),
        ("Результат. ", "Чек-лист законной публикации ресурса."),
        ("Презентация. ", "Тема 1. Основные понятия и классификация ЦОР.pptx"),
    ])
    doc.add_heading("Тема 2. EdTech и платформы", level=1)
    add_bullets(doc, [
        ("Старт. ", "Короткое голосование: какой платформой уже пользовались и почему."),
        ("Фокус. ", "Сравнение по критериям, а не по известности бренда."),
        ("Активность. ", "Группа выбирает 2 LMS и 1 конструктор под свой курс."),
        ("Результат. ", "Матрица из рабочей тетради и устное обоснование."),
        ("Презентация. ", "Тема 2. Современные тенденции в EdTech.pptx"),
    ])
    doc.add_heading("Тема 3. LMS Moodle", level=1)
    add_bullets(doc, [
        ("Старт. ", "Покажите пустой курс и спросите, чего в нем не хватает студенту."),
        ("Фокус. ", "Структура, роли, оценивание, путь прохождения."),
        ("Активность. ", "Сборка мини-курса: 3 раздела, тест, задание, форум."),
        ("Результат. ", "Опубликованный тестовый курс, пройденный от лица студента."),
        ("Презентация. ", "Тема 3. Системы управления обучением LMS.pptx"),
    ])
    doc.add_heading("Тема 4. Проектирование ИС", level=1)
    add_bullets(doc, [
        ("Старт. ", "Группа выбирает одну знакомую предметную область."),
        ("Фокус. ", "Обследование → задачи автоматизации → BPMN → UML."),
        ("Активность. ", "Схема процесса «как есть» и целевой сценарий."),
        ("Результат. ", "Комплект диаграмм и короткое обоснование границ системы."),
        ("Презентация. ", "Тема 4. Основы проектирования информационных систем.pptx"),
    ])
    doc.add_heading("Подготовка до пары", level=1)
    add_bullets(doc, [
        "Проверить вход в LMS и демонстрационный курс.",
        "Распечатать или выложить рабочую тетрадь.",
        "Открыть рубрику оценивания и показать ее студентам до работы.",
        "Сверить правовые формулировки с актуальными нормами и локальными актами.",
        "Держать запасной мини-кейс, если группа заканчивает раньше.",
    ], numbered=True)
    doc.add_heading("Что должно остаться после пары", level=1)
    add_bullets(doc, [
        "Продукт команды: таблица, схема или черновик курса.",
        "Самооценка по двум-трем критериям рубрики.",
        "Список вопросов, которые переносятся на следующее занятие.",
    ])
    doc.add_heading("Обратная связь", level=1)
    add_para(doc, "Формула комментария: что уже получилось → что улучшить в первую очередь → какой следующий шаг и к какому сроку.")
    add_callout(doc, "Сначала назовите сильную сторону. Затем одну, максимум две зоны роста. Длинный список замечаний группа не удерживает.")
    save_doc(doc, "Методический сценарий проведения занятий")


def build_workbook():
    doc = new_doc("Рабочая тетрадь")
    add_cover(doc, "Для практической работы", "Рабочая тетрадь", "Бланки заполняются на занятии и прикладываются к мини-проекту.")
    add_para(doc, "ФИО, группа и дата — на каждом занятии. Пустые клетки предназначены для рукописного или электронного заполнения.")
    add_form_table(doc, ["ФИО", "Группа", "Дата", "Тема занятия"])

    doc.add_heading("Шаблон 1. Паспорт ЦОР", level=1)
    add_para(doc, "Заполните до того, как выбирать картинки и платформу.")
    add_form_table(doc, [
        "Название ресурса",
        "Целевая аудитория",
        "Учебная цель",
        "Формируемые результаты",
        "Формат ресурса",
        "Что делает студент, а не только читает",
        "Критерий успеха",
        "Методические риски",
        "Технические риски",
        "Правовые риски",
    ])

    doc.add_heading("Шаблон 2. Матрица сравнения платформ", level=1)
    add_para(doc, "Шкала 1–5. Вес критерия задайте до выставления оценок. Впишите названия платформ в шапку.")
    add_data_table(
        doc,
        ["Критерий", "Вес", "Платформа A", "Платформа B", "Платформа C"],
        [
            ["Удобство преподавателя", "", "", "", ""],
            ["Удобство студента", "", "", "", ""],
            ["Интеграции", "", "", "", ""],
            ["Аналитика", "", "", "", ""],
            ["Безопасность и нормы РФ", "", "", "", ""],
            ["Стоимость владения", "", "", "", ""],
            ["Масштабируемость", "", "", "", ""],
            ["Итого с учетом веса", "", "", "", ""],
        ],
        [6.0, 1.8, 3.2, 3.2, 3.2],
        header=False,
        row_twips=640,
    )
    add_form_table(doc, ["Победитель", "Три основания выбора", "Риски до запуска", "Чем риски закрываем"])

    doc.add_heading("Шаблон 3. Карта курса в LMS", level=1)
    add_form_table(doc, ["Название курса", "Для кого курс", "Результат курса целиком"])
    add_data_table(
        doc,
        ["№", "Цель модуля", "Материал", "Действие", "Проверка", "Завершение"],
        [["1", "", "", "", "", ""], ["2", "", "", "", "", ""], ["3", "", "", "", "", ""], ["4", "", "", "", "", ""], ["5", "", "", "", "", ""]],
        [1.2, 3.4, 3.2, 3.2, 3.2, 3.2],
        header=False,
        row_twips=780,
    )

    doc.add_heading("Шаблон 4. Правовой чек-лист", level=1)
    add_para(doc, "Отметьте «да» только если пункт можно подтвердить ссылкой, настройкой или документом.")
    checks = [
        "У каждого заимствованного материала проверены права",
        "Лицензия и атрибуция указаны у материала",
        "В публикации нет чужого контента без основания",
        "Персональные данные собираются только под заявленную цель",
        "Есть правовое основание обработки ПДн",
        "Доступ к данным и оценкам разграничен по ролям",
        "Ресурс доступен обучающимся с ОВЗ",
        "Ресурс соответствует локальным актам организации",
    ]
    add_data_table(
        doc,
        ["№", "Проверка", "Да", "Нет", "Комментарий"],
        [[str(i), text, "", "", ""] for i, text in enumerate(checks, start=1)],
        [1.2, 8.4, 1.5, 1.5, 4.8],
        header=False,
        row_twips=700,
    )

    doc.add_heading("Шаблон 5. Обследование предметной области", level=1)
    add_form_table(doc, [
        "Границы области: что входит",
        "Что сознательно не входит",
        "Участники и роли",
        "Как процесс идет сейчас",
        "Где ошибки, очереди и дубли",
        "Что автоматизируем первым",
        "Ожидаемый эффект",
        "Ограничения и нормы",
    ])

    doc.add_heading("Шаблон 6. Паспорт диаграммы", level=1)
    add_para(doc, "Один бланк на каждую схему BPMN или UML.")
    add_form_table(doc, [
        "Тип диаграммы",
        "Название процесса или сценария",
        "Основные сущности и роли",
        "Бизнес-правила",
        "Исключения и альтернативы",
        "Какое требование закрывает схема",
        "Как поймем, что после внедрения стало лучше",
    ])
    save_doc(doc, "Рабочая тетрадь")


def build_assessment():
    doc = new_doc("Банк оценочных средств")
    add_cover(doc, "Текущий контроль и зачет", "Банк оценочных средств", "Вопросы для обсуждения, ключи для преподавателя и рубрика практических работ.")
    add_para(doc, "Вопросы можно использовать устно, в тесте LMS или как входной билет. Ключи вынесены отдельно, чтобы студенческую часть можно было выдать без ответов.")

    blocks = [
        ("Тема 1. ЦОР и правовые аспекты", [
            "Чем ЦОР отличается от файла, который просто выложили в сеть?",
            "Какой федеральный закон прямо регулирует обработку персональных данных?",
            "Что нужно проверить перед тем, как поставить в курс чужое изображение?",
            "Чем открытая лицензия отличается от ситуации, когда лицензии нет?",
            "Зачем при проектировании ЦОР сверяться с ФГОС?",
        ]),
        ("Тема 2. EdTech", [
            "Что такое адаптивное обучение на уровне конкретного курса?",
            "Назовите два уместных сценария ИИ и одну границу, которую нельзя отдавать модели.",
            "Какие критерии важны при выборе LMS, если бюджет уже не единственный аргумент?",
            "Почему пилот надежнее демонстрации платформы на вебинаре вендора?",
            "Как понять, что команда преподавателей действительно сможет работать на платформе?",
        ]),
        ("Тема 3. LMS Moodle", [
            "Из каких обязательных частей состоит базовая структура курса?",
            "Зачем ограничивать роли, если «проще выдать всем права преподавателя»?",
            "Для чего нужен банк вопросов, а не один фиксированный тест?",
            "Какие показатели в LMS помогают заметить трудности до конца модуля?",
            "Как проверить, что курс методически цельный?",
        ]),
        ("Тема 4. Проектирование ИС", [
            "Что входит в предпроектное обследование?",
            "Чем диаграмма BPMN отличается от диаграммы вариантов использования UML?",
            "Когда уместна итеративная модель жизненного цикла?",
            "Назовите три типичные ошибки раннего проектирования.",
            "Как связать диаграмму с требованием, чтобы схема не осталась картинкой?",
        ]),
    ]
    number = 1
    for title, questions in blocks:
        doc.add_heading(title, level=1)
        add_bullets(doc, [f"{number + i}. {q}" for i, q in enumerate(questions)], numbered=False)
        number += len(questions)

    doc.add_page_break()
    doc.add_heading("Ключи для преподавателя", level=1)
    add_callout(doc, "Этот раздел не выдается студентам вместе с вопросами. Формулировки короткие: на занятии ответ можно развернуть.")
    keys = [
        "Есть учебная цель, действие студента и способ проверки результата.",
        "Федеральный закон от 27.07.2006 № 152-ФЗ «О персональных данных».",
        "Правообладателя, вид лицензии и допустимые способы использования.",
        "Открытая лицензия задает условия. Отсутствие лицензии не разрешает использование.",
        "ФГОС задает ориентир результатов, с которыми связывают ресурс.",
        "Траектория и сложность подстраиваются под результат обучающегося.",
        "Например, черновик заданий и персональная подсказка. Граница: итоговую аттестацию подтверждает преподаватель.",
        "Цели, удобство, интеграции, аналитика, безопасность, полная стоимость владения.",
        "Пилот показывает реальную работу команды и скрытые ограничения.",
        "По скорости осваивания, качеству поддержки и устойчивости процесса без одного энтузиаста.",
        "Разделы, материалы, активности, контроль и условия завершения.",
        "Так защищают данные и не размывают ответственность за курс.",
        "Чтобы варианты различались, а проверка не превращалась в запоминание порядка ответов.",
        "Входы, прогресс по модулям, трудные вопросы, участие в обсуждении.",
        "Цель модуля, действие студента и критерий оценки говорят об одном и том же.",
        "Границы, роли, процесс как есть, проблемы и ограничения.",
        "BPMN описывает ход процесса. Use Case описывает цели действующих лиц на границе системы.",
        "Когда требования уточняются, а результат нужен частями.",
        "Нет проверки с пользователями, избыточная детализация, разрыв с целью и с реализацией.",
        "У схемы есть ссылка на требование и на критерий приемки.",
    ]
    add_bullets(doc, [f"{i}. {text}" for i, text in enumerate(keys, start=1)])

    doc.add_heading("Рубрика практической работы", level=1)
    add_para(doc, "Каждый критерий: 0, 1 или 2 балла. Максимум — 10.")
    add_data_table(
        doc,
        ["Критерий", "0 баллов", "1 балл", "2 балла"],
        [
            ["Методическая целостность", "Цель, действие и проверка не связаны", "Связь есть частично", "Цель, активность и контроль собраны в одну линию"],
            ["Структура", "Трудно понять порядок работы", "Порядок в целом ясен", "Структура понятна без устных пояснений"],
            ["Подача", "Текст перегружен или не читается", "Читается, но неровно", "Спокойно, контрастно, по делу"],
            ["Применимость", "Результат нельзя использовать", "Нужна заметная доработка", "Можно брать на занятие или в пилот"],
            ["Правовая корректность", "Есть нарушения", "Есть пробелы без критичного нарушения", "Нормы РФ и локальные правила учтены"],
        ],
        [3.6, 4.4, 4.4, 5.0],
    )
    add_para(doc, "0–4 балла — работу нужно существенно переделать. 5–7 — базовый рабочий уровень. 8–10 — материал готов к использованию на занятии.")

    doc.add_heading("Мини-проект к зачету", level=1)
    add_bullets(doc, [
        "Один учебный модуль в LMS.",
        "Паспорт ЦОР и сами материалы модуля.",
        "Комплект BPMN и UML по выбранному процессу.",
        "Заполненный правовой чек-лист и вывод по рискам.",
    ])
    add_callout(doc, "Защита: 7 минут выступления и 3 минуты вопросов. Студент показывает продукт, а не пересказывает определения.")
    save_doc(doc, "Банк оценочных средств")


def build_legal():
    doc = new_doc("Правовая памятка (РФ)")
    add_cover(
        doc,
        "Ориентир для занятий, не юридическое заключение",
        "Правовая памятка",
        "Авторское право, персональные данные, доступность и публикация учебного цифрового контента.",
    )
    add_callout(doc, "Перед публикацией курса сверяйте актуальные редакции нормативных актов и локальные акты своей образовательной организации. Памятка задает учебный маршрут проверки, а не заменяет юриста организации.")

    doc.add_heading("1. Нормы, которые стоит держать в поле зрения", level=1)
    add_data_table(
        doc,
        ["Акт", "Зачем он в курсе"],
        [
            ["Гражданский кодекс РФ, часть IV", "Авторские и смежные права, лицензии, использование произведений"],
            ["Федеральный закон от 29.12.2012 № 273-ФЗ", "Рамка образовательной деятельности и качества обучения"],
            ["Федеральный закон от 27.07.2006 № 149-ФЗ", "Общие правила работы с информацией и ее защиты"],
            ["Федеральный закон от 27.07.2006 № 152-ФЗ", "Обработка персональных данных обучающихся и работников"],
            ["Федеральный закон от 24.11.1995 № 181-ФЗ", "Доступность среды для инвалидов, в том числе в образовании"],
            ["Федеральный закон от 29.12.2010 № 436-ФЗ", "Возрастная маркировка, если контент предназначен детям"],
            ["Действующие требования СанПиН", "Режим экранной нагрузки и организация цифровых занятий"],
        ],
        [7.4, 10.0],
    )

    doc.add_heading("2. Авторское право на практике", level=1)
    add_para(doc, "Перед любым внешним материалом пройдите четыре шага.")
    add_bullets(doc, [
        "Определить правообладателя.",
        "Проверить лицензию и запреты: переработка, коммерческое использование, территория.",
        "Зафиксировать основание, по которому материал попадает в курс.",
        "Указать автора, источник и лицензию рядом с материалом.",
    ], numbered=True)
    add_para(doc, "Если контент делает подрядчик, в договоре заранее фиксируют, кому принадлежат исключительные права, какой объем передается и можно ли перерабатывать материал.")

    doc.add_heading("3. Персональные данные", level=1)
    add_para(doc, "Минимальный контур по 152-ФЗ для учебного курса:")
    add_bullets(doc, [
        "Цель обработки названа до сбора полей.",
        "Состав данных не шире этой цели.",
        "Есть правовое основание обработки.",
        "Назначены ответственные, доступы разделены по ролям.",
        "Есть организационные и технические меры защиты.",
        "Понятны срок хранения и порядок удаления.",
    ])
    add_callout(doc, "В LMS не публикуют лишние сведения в открытом разделе, ограничивают выгрузку ведомостей и проверяют, какие данные уходят во внешние сервисы.", fill="F0FDFA")

    doc.add_heading("4. Локальные документы организации", level=1)
    add_bullets(doc, [
        "Положение об электронной информационно-образовательной среде.",
        "Политика обработки персональных данных.",
        "Регламент использования LMS и связанных сервисов.",
        "Положение о текущем контроле и промежуточной аттестации.",
        "Требования к доступности учебных материалов.",
    ])

    doc.add_heading("5. Чек-лист перед публикацией", level=1)
    add_bullets(doc, [
        "Права на все материалы подтверждены.",
        "Условия лицензий соблюдены, атрибуция стоит на месте.",
        "Лишние персональные данные не собираются.",
        "Роли в LMS настроены, а не оставлены «по умолчанию».",
        "Содержание связано с целью и с ФГОС.",
        "Есть контраст, структура, субтитры или текстовая альтернатива.",
        "Учтены действующие санитарные требования к цифровым занятиям.",
        "Риски записаны в паспорте ресурса, а не только проговорены устно.",
    ], numbered=True)

    doc.add_heading("6. Как говорить об этом со студентами", level=1)
    add_bullets(doc, [
        "«Найти в интернете» не значит «можно вставить в курс».",
        "Персональные данные собирают под конкретную задачу.",
        "У каждого материала должен быть понятный статус прав.",
        "Хороший курс одновременно достигает учебной цели и не нарушает чужие права.",
    ])
    save_doc(doc, "Правовая памятка (РФ)")


def build_visual():
    doc = new_doc("Визуальный гайд")
    add_cover(doc, "Для слайдов студентов и преподавателя", "Визуальный гайд", "Короткие правила, чтобы презентация читалась с последнего ряда.")
    doc.add_heading("Базовый стиль", level=1)
    add_data_table(
        doc,
        ["Элемент", "Ориентир"],
        [
            ["Формат", "16:9"],
            ["Фон", "Светлый #F8FAFC или темный #0F172A на титуле"],
            ["Основной акцент", "#2563EB"],
            ["Второй акцент", "#0F766E"],
            ["Предупреждение", "#B45309"],
            ["Текст", "#0F172A на светлом фоне, белый на темном"],
            ["Шрифт", "Calibri. Если в организации принят Inter или Manrope, используйте его везде одинаково"],
        ],
        [5.2, 12.2],
    )
    doc.add_heading("Типографика", level=1)
    add_bullets(doc, [
        "Заголовок слайда: 28–36 pt, полужирный.",
        "Текст: 16–20 pt. Мельче 14 pt на проекторе уже плохо читается.",
        "На слайде не больше шести коротких тезисов.",
        "Одна мысль — один заголовок. Заголовок говорит, что делать или что запомнить.",
    ])
    doc.add_heading("Макет", level=1)
    add_bullets(doc, [
        "Поля не меньше 0,4 дюйма, важные смыслы не прижимать к краю.",
        "Таблицы — для сравнения, нумерованные шаги — для алгоритма, карточки — для равных тезисов.",
        "Каждый третий-четвертый слайд меняет композицию: после списка идет схема, кейс или таблица.",
        "Декор без смысла убирать. Линия и номер полезны, случайные иконки — нет.",
    ])
    doc.add_heading("Проверка за пять минут", level=1)
    add_bullets(doc, [
        "Смысл слайда понятен примерно за пять секунд?",
        "Можно ли убрать треть текста и ничего не потерять?",
        "Виден ли переход к следующему слайду?",
        "Таблица читается без увеличения?",
        "Есть ли слайд, на котором говорят студенты, а не только преподаватель?",
    ], numbered=True)
    save_doc(doc, "Визуальный гайд презентаций")


def main() -> None:
    deck_one()
    deck_two()
    deck_three()
    deck_four()
    build_overview()
    build_instructor()
    build_workbook()
    build_assessment()
    build_legal()
    build_visual()
    print("PPTX:")
    for path in sorted(PPTX_DIR.glob("*.pptx")):
        prs = Presentation(path)
        print(f"  {path.name}: {len(prs.slides)} слайдов")
    print("DOCX:")
    for path in sorted(DOCX_DIR.glob("*.docx")):
        doc = Document(path)
        print(f"  {path.name}: {len(doc.paragraphs)} абзацев, {len(doc.tables)} таблиц")


if __name__ == "__main__":
    main()

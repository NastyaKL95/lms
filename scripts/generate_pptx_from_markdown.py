#!/usr/bin/env python3
"""Convert prepared markdown presentations into .pptx files."""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Tuple

from pptx import Presentation


Heading = str
Bullet = Tuple[int, str]
Section = Tuple[Heading, List[Bullet]]


def clean_md_text(text: str) -> str:
    text = text.strip()
    text = text.replace("**", "")
    text = text.replace("`", "")
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", text)
    return text.strip()


def parse_markdown(path: Path) -> Tuple[str, List[Section]]:
    lines = path.read_text(encoding="utf-8").splitlines()

    title = path.stem
    sections: List[Section] = []
    current_heading: str | None = None
    current_items: List[Bullet] = []

    def flush_current() -> None:
        nonlocal current_heading, current_items
        if current_heading is not None:
            sections.append((current_heading, current_items))
        current_heading = None
        current_items = []

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("# "):
            title = clean_md_text(stripped[2:])
            continue

        if stripped.startswith("## ") or stripped.startswith("### "):
            flush_current()
            current_heading = clean_md_text(re.sub(r"^#{2,3}\s+", "", stripped))
            continue

        if current_heading is None:
            continue

        bullet_match = re.match(r"^(\s*)([-*])\s+(.*)$", line)
        number_match = re.match(r"^(\s*)\d+\.\s+(.*)$", line)

        if bullet_match:
            spaces, _, value = bullet_match.groups()
            level = max(0, len(spaces) // 2)
            current_items.append((min(level, 2), clean_md_text(value)))
        elif number_match:
            spaces, value = number_match.groups()
            level = max(0, len(spaces) // 2)
            current_items.append((min(level, 2), clean_md_text(value)))
        else:
            current_items.append((0, clean_md_text(stripped)))

    flush_current()
    return title, sections


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def chunk_items(items: List[Bullet], chunk_size: int = 7) -> List[List[Bullet]]:
    if not items:
        return [[]]
    return [items[i : i + chunk_size] for i in range(0, len(items), chunk_size)]


def add_section_slides(prs: Presentation, heading: str, items: List[Bullet]) -> None:
    for index, chunk in enumerate(chunk_items(items)):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading if index == 0 else f"{heading} (продолжение)"
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        if not chunk:
            text_frame.text = "Содержимое раздела"
            continue

        for item_index, (level, text) in enumerate(chunk):
            if item_index == 0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()
            paragraph.text = text
            paragraph.level = level


def convert(markdown_path: Path, output_path: Path) -> None:
    title, sections = parse_markdown(markdown_path)
    prs = Presentation()
    add_title_slide(prs, title, "Сформировано автоматически из markdown-структуры")

    for heading, items in sections:
        add_section_slides(prs, heading, items)

    prs.save(output_path)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    sources = [
        (
            root / "presentation-topic-1-detailed.md",
            root / "presentation-topic-1-detailed.pptx",
        ),
        (
            root / "presentation-topic-2-detailed.md",
            root / "presentation-topic-2-detailed.pptx",
        ),
    ]

    for source, output in sources:
        convert(source, output)
        print(f"created: {output}")


if __name__ == "__main__":
    main()
